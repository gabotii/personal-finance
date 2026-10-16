from __future__ import annotations

import csv
import io
from datetime import datetime
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.config import get_settings
from tests.conftest import categories_by_name, create_transaction


def _seed(client: TestClient) -> None:
    expense = categories_by_name(client)
    income = categories_by_name(client, "INCOME")
    create_transaction(
        client,
        type="INCOME",
        amount="1000000",
        description="Sueldo",
        category_id=income["Salario"]["id"],
        transaction_date="2026-09-01",
    )
    create_transaction(
        client,
        amount="45000",
        description="Supermercado",
        notes="Compra semanal",
        category_id=expense["Alimentación"]["id"],
        transaction_date="2026-09-10",
    )
    create_transaction(
        client,
        amount="12000",
        description='=HYPERLINK("http://evil")',
        category_id=expense["Transporte"]["id"],
        transaction_date="2026-08-20",
    )


def _parse_csv(content: bytes) -> list[list[str]]:
    assert content.startswith(b"\xef\xbb\xbf"), "CSV must start with a UTF-8 BOM for Excel"
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


def test_csv_export_is_valid(auth_client: TestClient) -> None:
    _seed(auth_client)
    response = auth_client.get("/api/v1/export/csv")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["content-disposition"].startswith('attachment; filename="movimientos_')
    rows = _parse_csv(response.content)
    assert rows[0] == ["Fecha", "Tipo", "Categoría", "Descripción", "Notas", "Monto", "Moneda"]
    assert rows[1:] == [
        ["2026-09-10", "Gasto", "Alimentación", "Supermercado", "Compra semanal", "45000", "CLP"],
        ["2026-09-01", "Ingreso", "Salario", "Sueldo", "", "1000000", "CLP"],
        ["2026-08-20", "Gasto", "Transporte", '\'=HYPERLINK("http://evil")', "", "12000", "CLP"],
    ]


def test_csv_export_applies_filters(auth_client: TestClient) -> None:
    _seed(auth_client)
    rows = _parse_csv(
        auth_client.get(
            "/api/v1/export/csv",
            params={"type": "EXPENSE", "date_from": "2026-09-01", "date_to": "2026-09-30"},
        ).content
    )
    assert [row[3] for row in rows[1:]] == ["Supermercado"]

    by_text = _parse_csv(auth_client.get("/api/v1/export/csv", params={"q": "sueldo"}).content)
    assert [row[3] for row in by_text[1:]] == ["Sueldo"]


def test_csv_export_streams_many_rows_in_batches(auth_client: TestClient) -> None:
    category_id = categories_by_name(auth_client)["Compras"]["id"]
    for index in range(30):
        create_transaction(
            auth_client,
            category_id=category_id,
            description=f"Item {index}",
            transaction_date=f"2026-09-{index % 28 + 1:02d}",
        )
    rows = _parse_csv(auth_client.get("/api/v1/export/csv").content)
    assert len(rows) == 31
    assert len({row[3] for row in rows[1:]}) == 30


def test_csv_export_rejects_invalid_filters(auth_client: TestClient) -> None:
    response = auth_client.get(
        "/api/v1/export/csv", params={"date_from": "2026-10-01", "date_to": "2026-01-01"}
    )
    assert response.status_code == 422


def test_excel_export_is_valid_and_correct(auth_client: TestClient) -> None:
    _seed(auth_client)
    response = auth_client.get("/api/v1/export/excel")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Movimientos", "Resumen", "Categorías"]

    sheet = workbook["Movimientos"]
    assert [cell.value for cell in sheet[1]] == [
        "Fecha",
        "Tipo",
        "Categoría",
        "Descripción",
        "Notas",
        "Monto",
    ]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:F4"
    assert sheet["A2"].value == datetime(2026, 9, 10)
    assert sheet["A2"].number_format == "DD-MM-YYYY"
    assert sheet["F2"].value == -45000
    assert sheet["F3"].value == 1000000
    assert '"$"#,##0' in sheet["F2"].number_format
    # Formula injection: stored as text, never as a formula.
    assert sheet["D4"].value == '=HYPERLINK("http://evil")'
    assert sheet["D4"].data_type == "s"
    # Totals rows.
    assert sheet["E6"].value == "Total ingresos"
    assert sheet["F6"].value == 1000000
    assert sheet["F7"].value == -57000
    assert sheet["F8"].value == 943000

    summary = workbook["Resumen"]
    values = {
        summary.cell(row=r, column=1).value: summary.cell(row=r, column=2).value
        for r in range(4, 11)
    }
    assert values["Movimientos"] == 3
    assert values["Ingresos"] == 1000000
    assert values["Gastos"] == 57000
    assert values["Balance"] == 943000
    assert Decimal(str(values["Tasa de ahorro"])) == Decimal("94.30")

    categories = workbook["Categorías"]
    rows = [[c.value for c in row] for row in categories.iter_rows(min_row=2, max_row=4)]
    assert rows == [
        ["Gasto", "Alimentación", 1, 45000, 78.95],
        ["Gasto", "Transporte", 1, 12000, 21.05],
        ["Ingreso", "Salario", 1, 1000000, 100],
    ]


def test_excel_export_applies_filters(auth_client: TestClient) -> None:
    _seed(auth_client)
    response = auth_client.get("/api/v1/export/excel", params={"type": "INCOME"})
    sheet = load_workbook(io.BytesIO(response.content))["Movimientos"]
    descriptions = [row[3] for row in sheet.iter_rows(min_row=2, max_row=3, values_only=True)]
    assert descriptions == ["Sueldo", None]


def test_excel_export_with_no_rows(auth_client: TestClient) -> None:
    response = auth_client.get("/api/v1/export/excel")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook["Movimientos"].max_row >= 1
    summary = workbook["Resumen"]
    assert summary["B10"].value == "Sin ingresos en el período"


def test_excel_export_limit(auth_client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    _seed(auth_client)
    monkeypatch.setattr(get_settings(), "excel_max_rows", 2)
    response = auth_client.get("/api/v1/export/excel")
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "EXPORT_TOO_LARGE"
