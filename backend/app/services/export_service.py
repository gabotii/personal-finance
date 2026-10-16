from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from datetime import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.core.enums import Currency, TransactionType
from app.core.errors import ValidationError
from app.core.money import ZERO, percentage
from app.models.transaction import Transaction
from app.models.user import User
from app.repositories.transaction_repository import TransactionRepository
from app.schemas.transaction import TransactionFilters
from app.services.transaction_service import validate_filters
from app.utils.dates import today_in
from app.utils.spreadsheet import is_formula_like, sanitize_csv_text

TYPE_LABELS = {TransactionType.INCOME: "Ingreso", TransactionType.EXPENSE: "Gasto"}
CSV_HEADERS = ("Fecha", "Tipo", "Categoría", "Descripción", "Notas", "Monto", "Moneda")
UTF8_BOM = "﻿"

MONEY_FORMATS: dict[Currency, str] = {
    Currency.CLP: '"$"#,##0',
    Currency.USD: '"US$"#,##0.00',
    Currency.EUR: '"€"#,##0.00',
}
DATE_FORMAT = "DD-MM-YYYY"
PERCENT_FORMAT = '0.00"%"'

HEADER_FILL = PatternFill("solid", start_color="1E293B", end_color="1E293B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="0F172A")
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", start_color="F1F5F9", end_color="F1F5F9")
INCOME_FONT = Font(color="047857")
EXPENSE_FONT = Font(color="B91C1C")
THIN_TOP = Border(top=Side(style="thin", color="94A3B8"))


def signed_money_format(currency: Currency) -> str:
    base = MONEY_FORMATS[currency]
    return f"{base};[Red]-{base}"


def export_filename(extension: str, user: User) -> str:
    return f"movimientos_{today_in(user.timezone).isoformat()}.{extension}"


def _format_amount(amount: Decimal, currency: Currency) -> str:
    return f"{amount:.{currency.decimal_places}f}"


class ExportService:
    def __init__(self, db: Session, settings: Settings) -> None:
        self.db = db
        self.settings = settings
        self.transactions = TransactionRepository(db)

    # ------------------------------------------------------------------- CSV

    def stream_csv(self, user: User, filters: TransactionFilters) -> Iterator[str]:
        """Yield the CSV in chunks; rows are fetched in keyset-paginated batches."""
        validate_filters(filters)
        currency = user.currency
        buffer = io.StringIO()
        writer = csv.writer(buffer, lineterminator="\r\n")

        def flush() -> str:
            chunk = buffer.getvalue()
            buffer.seek(0)
            buffer.truncate(0)
            return chunk

        writer.writerow(CSV_HEADERS)
        yield UTF8_BOM + flush()
        try:
            for batch in self.transactions.iter_batches(user.id, filters):
                for item in batch:
                    writer.writerow(
                        (
                            item.transaction_date.isoformat(),
                            TYPE_LABELS[item.type],
                            sanitize_csv_text(item.category.name),
                            sanitize_csv_text(item.description),
                            sanitize_csv_text(item.notes),
                            _format_amount(item.amount, currency),
                            currency.value,
                        )
                    )
                yield flush()
        finally:
            # The generator runs while the response streams; release the connection here.
            self.db.close()

    # ----------------------------------------------------------------- Excel

    def build_excel(self, user: User, filters: TransactionFilters) -> bytes:
        validate_filters(filters)
        total = self.transactions.count(user.id, filters)
        if total > self.settings.excel_max_rows:
            raise ValidationError(
                f"La exportación a Excel admite hasta {self.settings.excel_max_rows:,} "
                "movimientos. Acota los filtros o exporta a CSV.".replace(",", "."),
                code="EXPORT_TOO_LARGE",
                details={"total": total, "limit": self.settings.excel_max_rows},
            )
        rows = self.transactions.list_all(user.id, filters)

        workbook = Workbook()
        transactions_sheet = workbook.active
        if not isinstance(transactions_sheet, Worksheet):
            raise RuntimeError("A new workbook must have an active worksheet.")
        transactions_sheet.title = "Movimientos"
        self._write_transactions_sheet(transactions_sheet, rows, user.currency)
        self._write_summary_sheet(workbook.create_sheet("Resumen"), rows, user, filters)
        self._write_categories_sheet(workbook.create_sheet("Categorías"), rows, user.currency)

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _write_transactions_sheet(
        self, sheet: Worksheet, rows: list[Transaction], currency: Currency
    ) -> None:
        headers = ("Fecha", "Tipo", "Categoría", "Descripción", "Notas", "Monto")
        _write_header(sheet, 1, headers)
        money_format = signed_money_format(currency)

        income = expense = ZERO
        for index, item in enumerate(rows, start=2):
            signed = item.amount if item.type == TransactionType.INCOME else -item.amount
            if item.type == TransactionType.INCOME:
                income += item.amount
            else:
                expense += item.amount
            values: tuple[object, ...] = (
                item.transaction_date,
                TYPE_LABELS[item.type],
                item.category.name,
                item.description,
                item.notes or "",
                signed,
            )
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=index, column=column)
                _set_value(cell, value)
            sheet.cell(row=index, column=1).number_format = DATE_FORMAT
            sheet.cell(row=index, column=2).font = (
                INCOME_FONT if item.type == TransactionType.INCOME else EXPENSE_FONT
            )
            sheet.cell(row=index, column=6).number_format = money_format

        last_data_row = max(1, len(rows) + 1)
        sheet.auto_filter.ref = f"A1:F{last_data_row}"
        sheet.freeze_panes = "A2"

        totals_start = last_data_row + 2
        for offset, (label, value) in enumerate(
            (
                ("Total ingresos", income),
                ("Total gastos", -expense),
                ("Balance", income - expense),
            )
        ):
            row = totals_start + offset
            label_cell = sheet.cell(row=row, column=5, value=label)
            value_cell = sheet.cell(row=row, column=6, value=value)
            value_cell.number_format = money_format
            for cell in (label_cell, value_cell):
                cell.font = TOTAL_FONT
                cell.fill = TOTAL_FILL
            if offset == 0:
                label_cell.border = value_cell.border = THIN_TOP

        _set_widths(sheet, (13, 10, 22, 40, 36, 18))

    def _write_summary_sheet(
        self,
        sheet: Worksheet,
        rows: list[Transaction],
        user: User,
        filters: TransactionFilters,
    ) -> None:
        currency = user.currency
        money_format = signed_money_format(currency)
        income = sum((r.amount for r in rows if r.type == TransactionType.INCOME), ZERO)
        expense = sum((r.amount for r in rows if r.type == TransactionType.EXPENSE), ZERO)
        rate = percentage(income - expense, income)

        sheet["A1"] = "Resumen de movimientos"
        sheet["A1"].font = TITLE_FONT
        generated_at = datetime.now(ZoneInfo(user.timezone)).strftime("%d-%m-%Y %H:%M")
        sheet["A2"] = f"Generado el {generated_at} ({user.timezone})"
        sheet["A2"].font = Font(italic=True, color="64748B")

        period = _describe_period(filters)
        info: list[tuple[str, object, str | None]] = [
            ("Período", period, None),
            ("Moneda", currency.value, None),
            ("Movimientos", len(rows), "0"),
            ("Ingresos", income, money_format),
            ("Gastos", expense, money_format),
            ("Balance", income - expense, money_format),
            (
                "Tasa de ahorro",
                rate if rate is not None else "Sin ingresos en el período",
                PERCENT_FORMAT if rate is not None else None,
            ),
        ]
        for offset, (label, value, number_format) in enumerate(info):
            row = 4 + offset
            label_cell = sheet.cell(row=row, column=1, value=label)
            label_cell.font = TOTAL_FONT
            value_cell = sheet.cell(row=row, column=2)
            _set_value(value_cell, value)
            if number_format:
                value_cell.number_format = number_format
            value_cell.alignment = Alignment(horizontal="right")

        monthly: dict[str, dict[TransactionType, Decimal]] = {}
        for item in rows:
            key = item.transaction_date.strftime("%Y-%m")
            bucket = monthly.setdefault(
                key, {TransactionType.INCOME: ZERO, TransactionType.EXPENSE: ZERO}
            )
            bucket[item.type] += item.amount

        table_row = 4 + len(info) + 2
        sheet.cell(row=table_row - 1, column=1, value="Evolución mensual").font = TOTAL_FONT
        _write_header(sheet, table_row, ("Mes", "Ingresos", "Gastos", "Neto"))
        for offset, month in enumerate(sorted(monthly), start=1):
            month_income = monthly[month][TransactionType.INCOME]
            month_expense = monthly[month][TransactionType.EXPENSE]
            row = table_row + offset
            sheet.cell(row=row, column=1, value=month)
            for column, value in (
                (2, month_income),
                (3, month_expense),
                (4, month_income - month_expense),
            ):
                cell = sheet.cell(row=row, column=column, value=value)
                cell.number_format = money_format

        totals_row = table_row + len(monthly) + 1
        totals: tuple[object, ...] = ("Total", income, expense, income - expense)
        for column, total_value in enumerate(totals, start=1):
            total_cell = sheet.cell(row=totals_row, column=column)
            _set_value(total_cell, total_value)
            total_cell.font = TOTAL_FONT
            total_cell.fill = TOTAL_FILL
            total_cell.border = THIN_TOP
            if isinstance(total_value, Decimal):
                total_cell.number_format = money_format

        _set_widths(sheet, (24, 22, 18, 18))

    def _write_categories_sheet(
        self, sheet: Worksheet, rows: list[Transaction], currency: Currency
    ) -> None:
        money_format = MONEY_FORMATS[currency]
        totals: dict[tuple[TransactionType, str], tuple[Decimal, int]] = {}
        type_totals = {TransactionType.INCOME: ZERO, TransactionType.EXPENSE: ZERO}
        for item in rows:
            key = (item.type, item.category.name)
            amount, count = totals.get(key, (ZERO, 0))
            totals[key] = (amount + item.amount, count + 1)
            type_totals[item.type] += item.amount

        _write_header(sheet, 1, ("Tipo", "Categoría", "Movimientos", "Total", "% del tipo"))
        ordered = sorted(
            totals.items(),
            key=lambda entry: (entry[0][0] != TransactionType.EXPENSE, -entry[1][0], entry[0][1]),
        )
        for index, ((type_, name), (amount, count)) in enumerate(ordered, start=2):
            share = percentage(amount, type_totals[type_]) or ZERO
            values: tuple[object, ...] = (TYPE_LABELS[type_], name, count, amount, share)
            for column, value in enumerate(values, start=1):
                _set_value(sheet.cell(row=index, column=column), value)
            sheet.cell(row=index, column=1).font = (
                INCOME_FONT if type_ == TransactionType.INCOME else EXPENSE_FONT
            )
            sheet.cell(row=index, column=4).number_format = money_format
            sheet.cell(row=index, column=5).number_format = PERCENT_FORMAT

        last_data_row = max(1, len(ordered) + 1)
        sheet.auto_filter.ref = f"A1:E{last_data_row}"
        sheet.freeze_panes = "A2"

        row = last_data_row + 2
        for label, type_ in (
            ("Total gastos", TransactionType.EXPENSE),
            ("Total ingresos", TransactionType.INCOME),
        ):
            label_cell = sheet.cell(row=row, column=2, value=label)
            count_cell = sheet.cell(
                row=row, column=3, value=sum(c for (t, _), (_, c) in totals.items() if t == type_)
            )
            value_cell = sheet.cell(row=row, column=4, value=type_totals[type_])
            value_cell.number_format = money_format
            for cell in (label_cell, count_cell, value_cell):
                cell.font = TOTAL_FONT
                cell.fill = TOTAL_FILL
            row += 1

        _set_widths(sheet, (12, 26, 14, 18, 12))


def _set_value(cell: Cell | MergedCell, value: object) -> None:
    cell.value = value  # type: ignore[assignment]
    if isinstance(value, str) and is_formula_like(value):
        # openpyxl treats strings starting with "=" as formulas; force plain text.
        cell.data_type = "s"


def _write_header(sheet: Worksheet, row: int, headers: tuple[str, ...]) -> None:
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[row].height = 22


def _set_widths(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _describe_period(filters: TransactionFilters) -> str:
    if filters.date_from and filters.date_to:
        return f"{filters.date_from:%d-%m-%Y} a {filters.date_to:%d-%m-%Y}"
    if filters.date_from:
        return f"Desde {filters.date_from:%d-%m-%Y}"
    if filters.date_to:
        return f"Hasta {filters.date_to:%d-%m-%Y}"
    return "Todos los movimientos"
